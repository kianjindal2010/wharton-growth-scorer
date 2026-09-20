from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
PALE_BLUE = "EAF3F8"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_font(run, name: str = "Arial", size: float = 10.5, bold: bool = False, color: str = BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def borders(cell):
    props = cell._tc.get_or_add_tcPr()
    tc_borders = props.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        props.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = tc_borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LIGHT_GRAY)


def margins(cell, value: int = 110):
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def add_table(doc, headers, rows, widths, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        cell.text = header
        cell.width = Inches(width)
    for row in rows:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.text = str(value)
            cell.width = Inches(width)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            borders(cell)
            if row_index == 0:
                shade(cell, NAVY)
            elif row_index % 2 == 0:
                shade(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_font(run, size=font_size, bold=row_index == 0, color=WHITE if row_index == 0 else BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, lead: str, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(lead)
    run.bold = True
    paragraph.add_run(text)


def add_code(doc, text: str, size: float = 8.5):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_font(run, "Consolas", size)


def build(path: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    title_props = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_props.find(qn("w:pBdr"))
    if title_border is not None:
        title_props.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.add_run("Wharton Growth Scorer Results and User Guide")
    subtitle = doc.add_paragraph()
    set_font(subtitle.add_run("Version 0.4.1 for the equity growth sleeve"), size=13, bold=True)
    metadata = doc.add_paragraph()
    set_font(metadata.add_run("Prepared for the Wharton Global High School Investment Competition team\n"), size=9.5, bold=True)
    set_font(metadata.add_run("Repository  github.com/kianjindal2010/wharton-growth-scorer"), size=9.5)

    doc.add_paragraph(
        "The model gives the team a consistent way to screen an analyst-vetted stock, compare it with similar companies, and document the reasoning behind an investment decision. It combines financial quality, growth, valuation, balance-sheet strength, market behavior, sector-specific measures, and recent news. The result is a transparent 0 to 100 score, not an automatic trade instruction or a guaranteed return forecast."
    )

    doc.add_heading("Results in brief", level=1)
    add_table(
        doc,
        ["Evaluation", "Original model", "Sector aware model"],
        [
            ["Score versus six month USD return", "0.48 rank correlation", "0.57 rank correlation"],
            ["Score versus benchmark relative return", "0.35 rank correlation", "0.44 rank correlation"],
            ["Directional hit rate", "77.8 percent", "88.9 percent"],
            ["Top minus bottom score group return", "17.4 percentage points", "16.7 percentage points"],
        ],
        [3.05, 1.75, 1.75],
        8.9,
    )
    doc.add_paragraph(
        "The test used 18 companies across the United States, Japan, the United Kingdom, India, and Taiwan from 20 March to 18 September 2026. Thirteen companies had sufficient data for an actionable verdict. The results support using the model as a structured ranking and research tool, but the sample is too small and the period too short to prove forecasting ability."
    )
    doc.add_paragraph(
        "A separate six-stock semiconductor diagnostic tested Micron, SK Hynix, Samsung, Nanya, Winbond, and TSMC. It showed that memory producers need a cycle-aware scorecard that recognizes DRAM pricing, margin inflection, inventories, and shorter-term momentum. That redesign was informed by the same outcome and therefore still requires walk-forward testing."
    )

    doc.add_heading("What the model does", level=1)
    add_table(
        doc,
        ["Stage", "Model action", "Output"],
        [
            ["Collect", "Downloads prices, statements, exchange rates, benchmarks, corporate actions, and eligible Yahoo Finance news.", "Frozen dated snapshot"],
            ["Measure", "Calculates financial, valuation, risk, momentum, sector, and news variables without using data after the selected date.", "Comparable metrics"],
            ["Score", "Converts each metric to 0 to 100 using fixed bad, neutral, and excellent anchors, then applies weights and risk gates.", "Overall score and verdict"],
            ["Document", "Stores every input, source, warning, transformed factor, and contribution.", "Excel, JSON, and history record"],
        ],
        [0.75, 4.1, 1.7],
        8.7,
    )

    doc.add_page_break()
    doc.add_heading("Variables considered", level=1)
    doc.add_paragraph("The selected scorecard determines which variables receive weight. Missing metrics receive a neutral score of 50 and reduce confidence; their weights are not redistributed.")
    variable_rows = [
        ["Market data", "Open, high, low, close, adjusted close, volume, dividends, splits, FX rates, local benchmark prices"],
        ["Risk and momentum", "USD volatility, maximum drawdown, downside beta, and relative returns over 3, 6, and 12 months"],
        ["Quality and cash flow", "ROIC, ROE, ROA, gross and operating margins, FCF margin, cash conversion, and margin stability"],
        ["Growth", "Revenue, EPS, FCF and book-value growth; latest growth; acceleration; margin and inventory trends"],
        ["Valuation", "FCF, EBIT, earnings and shareholder yields; price to sales, book, and tangible book"],
        ["Financial strength", "Net debt to EBITDA, interest coverage, liquidity, cash to debt, debt to equity, and asset turnover"],
        ["Investment and innovation", "R&D intensity and growth, capex intensity, incremental ROIC, and revenue growth relative to capex"],
        ["Specialist factors", "Bank capital and asset quality, insurer solvency and reserves, biotechnology runway and dilution, and DRAM contract pricing"],
        ["News", "30-day sentiment, 30-versus-90-day trend, coverage quality, relevance, source quality, and recency"],
    ]
    add_table(doc, ["Variable group", "Examples"], variable_rows, [1.55, 5.0], 8.7)

    doc.add_heading("Sector aware scorecards", level=1)
    doc.add_paragraph(
        "The automatic classifier selects among general, technology, healthcare, financial platform, industrial, consumer, energy and materials, bank, insurer, pre-profit biotechnology, semiconductor, and memory-semiconductor scorecards. This prevents one set of assumptions from being applied to economically different businesses. Core financial and market factors contribute 95 percent of the score; the common news layer contributes 5 percent."
    )

    doc.add_heading("How the data is sourced", level=1)
    source_rows = [
        ["Yahoo Finance", "Adjusted OHLCV, dividends, splits, available annual and quarterly statements, company classification, benchmarks, FX rates, and recent news"],
        ["Verified overrides", "Regulatory or company filings for metrics not reliably supplied by Yahoo Finance, including CET1, solvency, reserves, and specialist cycle data"],
        ["Analyst controls", "Exact ticker, country, as-of date, manual scorecard override when justified, source URL, publication date, and verifier"],
    ]
    add_table(doc, ["Source", "Use"], source_rows, [1.45, 5.1], 8.8)
    doc.add_paragraph(
        "Every run applies an information cutoff. Prices and statement periods after the as-of date are excluded. News must have a publication timestamp on or before that date, mention the company or ticker, and survive duplicate filtering. If historical Yahoo Finance news is unavailable, news remains neutral instead of using current headlines."
    )

    doc.add_heading("Risk controls", level=1)
    add_bullet(doc, "Confidence  ", "Below 70 percent produces Insufficient Data.")
    add_bullet(doc, "Freshness  ", "Prices more than five trading days stale or statements more than 21 months old produce Insufficient Data.")
    add_bullet(doc, "Risk gates  ", "Negative equity, severe leverage, inadequate regulatory capital, inadequate solvency, or biotechnology runway below 12 months can override the numerical score.")

    doc.add_heading("Installation", level=1)
    doc.add_paragraph(
        "The repository and installer are public. No GitHub account or Git installation is required. The teammate needs Python 3.11 or 3.12 installed on Windows with the Add Python to PATH option selected."
    )
    doc.add_heading("One line PowerShell installation", level=2)
    doc.add_paragraph("Paste the complete command below into PowerShell:")
    one_line = "irm https://raw.githubusercontent.com/kianjindal2010/wharton-growth-scorer/main/bootstrap.ps1 | iex"
    add_code(doc, one_line, 8.0)
    doc.add_paragraph("The same line updates an existing installation. After installation, the command can be run from PowerShell or Command Prompt.")

    doc.add_heading("Run the model", level=1)
    add_code(doc, "wharton predict", 10)
    doc.add_paragraph("The interactive program asks for the following inputs:")
    prompt_rows = [
        ["Stock ticker", "Exact Yahoo Finance format", "MSFT, 7203.T, AZN.L, TCS.NS, 2330.TW, 000660.KS"],
        ["Country", "Matching market code", "US, JP, GB, IN, TW, KR"],
        ["As-of date", "YYYY-MM-DD or Enter for today", "2026-09-20"],
        ["Scorecard", "Press Enter for automatic selection", "auto"],
        ["Override", "Press Enter unless a verified specialist file exists", "Optional XLSX or CSV path"],
    ]
    add_table(doc, ["Prompt", "Required format", "Example"], prompt_rows, [1.25, 2.45, 2.85], 8.6)

    doc.add_heading("Direct command example", level=2)
    add_code(doc, "wharton predict --ticker 2330.TW --country TW --as-of 2026-09-20 --scorecard auto", 9.0)

    doc.add_page_break()
    doc.add_heading("Excel output", level=1)
    doc.add_paragraph("The command prints the exact file location and saves the workbook under:")
    add_code(doc, "Documents\\Wharton Growth Scorer\\output\\scores\\YYYY-MM-DD\\", 9.0)
    output_rows = [
        ["Summary", "Overall score, verdict, confidence, scorecard, rank, risk gates, and pillar scores"],
        ["Factor Breakdown", "Every variable, raw value, anchor transformation, weight, contribution, observation status, and source"],
        ["Raw Data", "Prices, statements, FX data, benchmark data, news records, and calculated metrics"],
        ["Source Log", "Source, retrieval time, reporting period, URL, and override provenance"],
        ["Warnings", "Missing information, stale fields, clipping, classification concerns, and triggered gates"],
    ]
    add_table(doc, ["Worksheet", "Contents"], output_rows, [1.55, 5.0], 8.8)

    doc.add_heading("How to use the result", level=1)
    doc.add_paragraph(
        "A score of 75 or more is a Buy Candidate, 60 to 74.99 is Watch, and below 60 is Reject, subject to confidence and risk gates. Analysts should compare companies within the same scorecard and weekly cycle, read the warnings and factor breakdown, and use the score to organize research rather than replace judgment. The model does not confirm WInS eligibility, size positions, or manage the defensive sleeve."
    )

    doc.add_heading("Key references", level=1)
    doc.add_paragraph("Repository and installation files  https://github.com/kianjindal2010/wharton-growth-scorer")
    doc.add_paragraph("Complete variable dictionary  VARIABLES.md in the repository")
    doc.add_paragraph("Detailed backtest  reports/Growth_Sleeve_Six_Month_Backtest_Report.docx in the repository")

    doc.core_properties.title = "Wharton Growth Scorer Results and User Guide"
    doc.core_properties.subject = "Concise model overview, results, variables, data sources, installation, and operation"
    doc.core_properties.author = "Wharton Global High School Investment Competition Team"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build(Path("reports/Wharton_Growth_Scorer_Results_and_User_Guide.docx"))
