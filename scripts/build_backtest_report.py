from __future__ import annotations

import json
import math
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
BLUE = "2079B0"
PALE_BLUE = "EAF3F8"
LIGHT_GRAY = "D9D9D9"
RED = "BE1E2D"
GREEN = "2E7D32"
WHITE = "FFFFFF"
BLACK = "000000"


def pct(value):
    return "N/A" if value is None else f"{value * 100:.1f}%"


def num(value):
    return "N/A" if value is None else f"{value:.2f}"


def font(size, bold=False):
    path = Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size=size)


def scatter_chart(records, path: Path):
    width, height = 1400, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, axis_font, label_font = font(34, True), font(22), font(18)
    draw.text((70, 28), "Model score compared with subsequent six month USD return", fill="#000000", font=title_font)
    left, top, right, bottom = 105, 105, width - 55, height - 100
    x_min, x_max = 35.0, 85.0
    returns = [record["usd_return"] * 100 for record in records if record["usd_return"] is not None]
    y_min = math.floor((min(returns) - 5) / 10) * 10
    y_max = math.ceil((max(returns) + 5) / 10) * 10
    draw.line((left, bottom, right, bottom), fill="#333333", width=3)
    draw.line((left, top, left, bottom), fill="#333333", width=3)
    for x in range(40, 86, 10):
        px = left + (x - x_min) / (x_max - x_min) * (right - left)
        draw.line((px, top, px, bottom), fill="#E5E5E5", width=1)
        draw.text((px - 12, bottom + 15), str(x), fill="#333333", font=axis_font)
    step = 10 if y_max - y_min <= 80 else 20
    for y in range(int(y_min), int(y_max) + 1, step):
        py = bottom - (y - y_min) / (y_max - y_min) * (bottom - top)
        draw.line((left, py, right, py), fill="#E5E5E5", width=1)
        draw.text((25, py - 12), f"{y}%", fill="#333333", font=axis_font)
    colors = {"Buy Candidate": "#2E7D32", "Watch": "#2079B0", "Reject": "#BE1E2D", "Insufficient Data": "#777777"}
    for record in records:
        if record["usd_return"] is None:
            continue
        px = left + (record["score"] - x_min) / (x_max - x_min) * (right - left)
        py = bottom - (record["usd_return"] * 100 - y_min) / (y_max - y_min) * (bottom - top)
        color = colors[record["verdict"]]
        draw.ellipse((px - 8, py - 8, px + 8, py + 8), fill=color, outline="white", width=2)
        draw.text((px + 10, py - 12), record["ticker"], fill="#222222", font=label_font)
    draw.text(((left + right) // 2 - 60, height - 52), "Model score", fill="#000000", font=axis_font)
    draw.text((8, 72), "USD return", fill="#000000", font=axis_font)
    legend_x = width - 450
    for index, verdict in enumerate(["Buy Candidate", "Watch", "Reject", "Insufficient Data"]):
        x = legend_x + (index % 2) * 220
        y = 55 + (index // 2) * 32
        draw.ellipse((x, y, x + 14, y + 14), fill=colors[verdict])
        draw.text((x + 22, y - 5), verdict, fill="#222222", font=label_font)
    image.save(path)


def bucket_chart(summary, path: Path):
    values = [
        ("Buy Candidate", summary["buy_mean_usd_return"], f"n={summary['buy_candidates']}", "#2E7D32"),
        ("Watch", summary["watch_mean_usd_return"], f"n={summary['watch_companies']}", "#2079B0"),
        ("Reject", summary["reject_mean_usd_return"], f"n={summary['reject_companies']}", "#BE1E2D"),
    ]
    width, height = 1200, 650
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 30), "Average six month USD return by model verdict", fill="#000000", font=font(34, True))
    left, top, right, bottom = 120, 125, width - 70, height - 105
    draw.line((left, bottom, right, bottom), fill="#333333", width=3)
    draw.line((left, top, left, bottom), fill="#333333", width=3)
    max_y = 0.40
    for value in range(0, 41, 10):
        py = bottom - (value / 100) / max_y * (bottom - top)
        draw.line((left, py, right, py), fill="#E5E5E5", width=1)
        draw.text((45, py - 12), f"{value}%", fill="#333333", font=font(20))
    bar_width = 180
    centers = [300, 600, 900]
    for (label, value, sample, color), center in zip(values, centers):
        height_px = value / max_y * (bottom - top)
        draw.rectangle((center - bar_width // 2, bottom - height_px, center + bar_width // 2, bottom), fill=color)
        draw.text((center - 42, bottom - height_px - 38), pct(value), fill="#000000", font=font(24, True))
        label_width = draw.textlength(label, font=font(22, True))
        draw.text((center - label_width / 2, bottom + 16), label, fill="#000000", font=font(22, True))
        draw.text((center - 18, bottom + 49), sample, fill="#555555", font=font(19))
    image.save(path)


def semiconductor_comparison_chart(initial_records, enhanced_records, path: Path):
    initial = {row["ticker"]: row for row in initial_records}
    width, height = 1400, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 28), "Semiconductor scorecard diagnostic: original and cycle-aware scores", fill="#000000", font=font(32, True))
    left, top, right, bottom = 110, 115, width - 70, height - 125
    draw.line((left, bottom, right, bottom), fill="#333333", width=3)
    draw.line((left, top, left, bottom), fill="#333333", width=3)
    for value in range(40, 91, 10):
        py = bottom - (value - 40) / 50 * (bottom - top)
        draw.line((left, py, right, py), fill="#E5E5E5", width=1)
        draw.text((45, py - 12), str(value), fill="#333333", font=font(20))
    tickers = [row["ticker"] for row in enhanced_records]
    spacing = (right - left) / len(tickers)
    for index, row in enumerate(enhanced_records):
        center = left + spacing * (index + 0.5)
        old_score = initial[row["ticker"]]["score"]
        new_score = row["score"]
        for offset, value, color in ((-20, old_score, "#9E9E9E"), (20, new_score, "#2079B0")):
            py = bottom - (value - 40) / 50 * (bottom - top)
            draw.rectangle((center + offset - 15, py, center + offset + 15, bottom), fill=color)
            draw.text((center + offset - 24, py - 29), f"{value:.0f}", fill="#222222", font=font(17, True))
        label_width = draw.textlength(row["ticker"], font=font(18, True))
        draw.text((center - label_width / 2, bottom + 18), row["ticker"], fill="#222222", font=font(18, True))
    draw.rectangle((width - 380, 62, width - 356, 86), fill="#9E9E9E")
    draw.text((width - 345, 59), "Original semiconductor", fill="#222222", font=font(18))
    draw.rectangle((width - 380, 91, width - 356, 115), fill="#2079B0")
    draw.text((width - 345, 88), "Cycle-aware memory", fill="#222222", font=font(18))
    image.save(path)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths=None, font_size=8.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[column_index])
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
    style_table(table, widths, font_size)
    return table


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def build_report(payload_path: Path, enhanced_path: Path, semiconductor_path: Path, initial_semiconductor_path: Path, output_path: Path, chart_dir: Path):
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    enhanced = json.loads(enhanced_path.read_text(encoding="utf-8"))
    semiconductor = json.loads(semiconductor_path.read_text(encoding="utf-8"))
    initial_semiconductor = json.loads(initial_semiconductor_path.read_text(encoding="utf-8"))
    summary = payload["summary"]
    records = payload["records"]
    chart_dir.mkdir(parents=True, exist_ok=True)
    scatter_path = chart_dir / "score_vs_return.png"
    bucket_path = chart_dir / "return_by_verdict.png"
    semiconductor_chart_path = chart_dir / "semiconductor_scorecard_comparison.png"
    scatter_chart(records, scatter_path)
    bucket_chart(summary, bucket_path)
    semiconductor_comparison_chart(initial_semiconductor["records"], semiconductor["records"], semiconductor_chart_path)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 27), ("Heading 1", 18), ("Heading 2", 13)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(BLACK)
        styles[name].font.bold = True
        styles[name].paragraph_format.space_before = Pt(12)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.keep_with_next = True
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.add_run("Growth Sleeve Scoring Model Backtest and Semiconductor Cycle Analysis")
    title.paragraph_format.space_after = Pt(10)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("Retrospective screening test across 18 companies, plus a six-stock DRAM and foundry diagnostic")
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.bold = True
    metadata = doc.add_paragraph()
    metadata.add_run("Scoring date 20 March 2026\n").bold = True
    metadata.add_run("Performance comparison through 18 September 2026\n")
    metadata.add_run("Prepared for the Wharton Global High School Investment Competition team")

    doc.add_paragraph(
        "The model showed useful ranking information in this limited test, but the result is not yet strong enough to treat the score as a return forecast. "
        f"Across {summary['actionable_companies']} companies with sufficient data, score rank had a {summary['spearman_score_vs_usd_return']:.2f} correlation with subsequent USD return. "
        f"The highest-score quartile returned {pct(summary['top_quartile_mean_usd_return'])} on average compared with {pct(summary['bottom_quartile_mean_usd_return'])} for the lowest quartile, a spread of {pct(summary['top_minus_bottom_spread'])}. "
        "The evidence supports using the model as a disciplined screening and comparison tool while expanding the test period and improving point-in-time filing data."
    )

    doc.add_heading("Main findings", level=1)
    finding_rows = [
        ("Coverage", f"{summary['actionable_companies']} of {summary['companies']} companies produced actionable verdicts; {summary['insufficient_data_companies']} were marked Insufficient Data."),
        ("Ranking signal", f"Spearman rank correlation was {summary['spearman_score_vs_usd_return']:.2f} against USD return and {summary['spearman_score_vs_usd_excess_return']:.2f} against local-benchmark-relative USD return."),
        ("Score separation", f"Top quartile minus bottom quartile USD return was {pct(summary['top_minus_bottom_spread'])}."),
        ("Verdict groups", f"Watch names returned {pct(summary['watch_mean_usd_return'])} on average; Reject names returned {pct(summary['reject_mean_usd_return'])}."),
        ("Buy result", f"The single Buy Candidate returned {pct(summary['buy_mean_usd_return'])} in USD but {pct(summary['buy_mean_usd_excess_return'])} relative to its local benchmark."),
        ("Directional test", f"The directional hit rate was {pct(summary['directional_hit_rate'])} for Buy outperformance and Reject underperformance. Watch names were excluded from this binary test."),
    ]
    add_table(doc, ["Measure", "Result"], finding_rows, widths=[1.35, 5.45], font_size=9.5)

    doc.add_page_break()
    doc.add_heading("All-sector enhancement result", level=1)
    enhanced_summary = enhanced["summary"]
    all_sector_rows = [
        ("Score versus USD return", num(summary["spearman_score_vs_usd_return"]), num(enhanced_summary["spearman_score_vs_usd_return"])),
        ("Score versus benchmark-relative return", num(summary["spearman_score_vs_usd_excess_return"]), num(enhanced_summary["spearman_score_vs_usd_excess_return"])),
        ("Top-minus-bottom USD return spread", pct(summary["top_minus_bottom_spread"]), pct(enhanced_summary["top_minus_bottom_spread"])),
        ("Directional hit rate", pct(summary["directional_hit_rate"]), pct(enhanced_summary["directional_hit_rate"])),
    ]
    add_table(doc, ["Measure", "Original", "Sector-aware"], all_sector_rows, widths=[3.25, 1.55, 1.55], font_size=9.2)
    doc.add_paragraph(
        "Across the same 18-company universe, sector-aware scorecards improved score/return rank correlation from 0.48 to 0.57, benchmark-relative rank correlation from 0.35 to 0.44, and directional hit rate from 77.8% to 88.9%. The top-minus-bottom spread remained positive but declined slightly from 17.4 to 16.7 percentage points. This is a useful diagnostic, not independent proof: the sector design was developed after reviewing the first result."
    )

    doc.add_heading("Test design", level=1)
    doc.add_paragraph(
        "Each company was scored using only prices and financial-statement periods dated on or before 20 March 2026. The model selected the scorecard automatically, applied its fixed metric anchors and weights, calculated data confidence, and then applied the prescribed risk gates. No WInS eligibility decision was made."
    )
    doc.add_paragraph(
        "Subsequent performance was measured from the first available adjusted close on or after the scoring date through 18 September 2026. Adjusted prices include the effect of splits and distributions. Local returns were converted to US dollars using the model's market-specific FX series, and each company was also compared with its local market benchmark."
    )
    design_rows = [
        ("Universe", "18 companies across US, Japan, United Kingdom, India, and Taiwan"),
        ("Scorecards", "General company, bank, and insurer scorecards selected automatically"),
        ("Decision labels", "Buy Candidate 75 or above; Watch 60 to 74.99; Reject below 60; Insufficient Data below 70% confidence or stale inputs"),
        ("Return basis", "Adjusted-close return in local currency and US dollars; excess return versus local benchmark"),
        ("Data source", "Yahoo Finance prices, corporate actions, FX rates, benchmark prices, and available statements"),
    ]
    add_table(doc, ["Element", "Method"], design_rows, widths=[1.35, 5.45], font_size=9.2)

    doc.add_heading("Sector-aware architecture", level=1)
    sector_rows = [
        ("Technology", "Gross margin, R&D intensity, scalable FCF, growth-adjusted valuation"),
        ("Healthcare", "Product economics, innovation spending, growth, resilience; separate pre-profit biotech rules"),
        ("Financial services", "Separate bank, insurer, and financial-platform scorecards"),
        ("Industrials", "ROIC, asset turnover, capex intensity, inventory cycle, leverage"),
        ("Consumer/media", "Brand economics, margins, inventory turnover, cash conversion"),
        ("Energy/materials/utilities", "Cash returns, balance sheet, capital discipline, cyclicality"),
    ]
    add_table(doc, ["Team mandate", "Specialized emphasis"], sector_rows, widths=[1.8, 4.95], font_size=9.0)

    doc.add_heading("News sentiment layer", level=1)
    doc.add_paragraph(
        "Every scorecard now reserves 5% for news: weighted 30-day sentiment (3%), the 30-day versus 90-day trend (1%), and quality-adjusted coverage (1%). Yahoo Finance headlines and summaries are deduplicated and must mention the company or ticker. Each item is then weighted by company relevance, publisher quality, and a 14-day recency half-life. A finance-aware deterministic lexicon handles earnings beats and misses, guidance changes, approvals, investigations, recalls, downgrades, dilution and other market events, including basic negation and intensifiers."
    )
    doc.add_paragraph(
        "Publication timestamps are a hard point-in-time boundary. Yahoo Finance did not expose March 2026 headlines during this September 2026 retrospective run, so all historical news factors stayed neutral and reduced confidence by five percentage points. Therefore, the all-sector backtest improvement above comes from sector specialization—not from sentiment. Current-run news ingestion was separately validated, but historical sentiment requires an archived news dataset before it can be backtested honestly."
    )

    doc.add_heading("Company results", level=1)
    doc.add_paragraph("Results are ordered from highest to lowest model score. Insufficient Data observations are shown but excluded from correlation, quartile, and verdict-group performance calculations.")
    ordered = sorted(records, key=lambda record: record["score"], reverse=True)
    result_rows = [
        [
            record["ticker"],
            record["country"],
            f"{record['score']:.1f}",
            record["verdict"].replace(" Candidate", ""),
            f"{record['confidence']:.0f}%",
            pct(record["usd_return"]),
            pct(record["usd_excess_return"]),
        ]
        for record in ordered
    ]
    add_table(
        doc,
        ["Ticker", "Market", "Score", "Verdict", "Confidence", "USD return", "USD excess"],
        result_rows,
        widths=[0.82, 0.55, 0.58, 1.08, 0.75, 0.82, 0.82],
        font_size=8.0,
    )

    doc.add_heading("Score and return relationship", level=1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(scatter_path), width=Inches(6.75))
    caption = doc.add_paragraph("Figure 1  Model score and subsequent USD return. Gray points lacked sufficient data for an actionable verdict.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].font.italic = True
    caption.runs[0].font.size = Pt(9)

    doc.add_page_break()
    doc.add_heading("Performance by verdict", level=1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(bucket_path), width=Inches(6.55))
    caption = doc.add_paragraph("Figure 2  Average USD return for actionable verdict groups. The Buy Candidate group contains one company and must not be generalized.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].font.italic = True
    caption.runs[0].font.size = Pt(9)

    doc.add_heading("What the results suggest", level=1)
    doc.add_paragraph(
        "The clearest positive result is separation between the higher- and lower-scoring groups. Watch companies averaged materially higher returns than Reject companies, and the top-quartile spread was positive. This supports the scorecard's intended use as a consistent ranking framework."
    )
    doc.add_paragraph(
        "The model was less successful as a benchmark-relative signal. Taiwan Semiconductor Manufacturing was the only Buy Candidate and gained strongly in US-dollar terms, yet it underperformed the Taiwan benchmark over the same window. TCS also received a relatively high Watch score but declined. Conversely, Caterpillar and Hon Hai produced positive absolute returns despite Reject verdicts. These cases show why the score should organize research rather than replace portfolio judgment."
    )
    doc.add_paragraph(
        f"Specialist coverage remains a material weakness. {summary['insufficient_data_companies']} financial or insurance names were marked Insufficient Data because Yahoo Finance did not provide all required specialist metrics. This is the desired safety behavior, but it also means verified CET1, efficiency, asset-quality, solvency, leverage, and reserve data must be entered before the model can compare those companies confidently."
    )

    doc.add_page_break()
    doc.add_heading("DRAM and TSMC diagnostic", level=1)
    doc.add_paragraph(
        "The requested DRAM test uses five listed memory manufacturers—Micron, SK Hynix, Samsung Electronics, Nanya Technology, and Winbond Electronics—rather than the DRAM exchange-traded fund. The ETF was not available at the 20 March 2026 scoring date and the model is intentionally stock-only. Taiwan Semiconductor Manufacturing (TSMC) is included as a foundry comparator, not as a DRAM producer."
    )
    doc.add_paragraph(
        "The first run applied one long-term semiconductor scorecard to all six companies. It failed at this cyclical turning point: three memory producers were rejected even though they subsequently rallied sharply. The diagnosis was structural, not cosmetic. Trailing earnings and free cash flow remain depressed near a memory-cycle trough, while contract pricing, revenue acceleration, margin inflection, inventories, and shorter-horizon momentum can improve first."
    )
    original_summary = initial_semiconductor["summary"]
    enhanced_summary = semiconductor["summary"]
    comparison_rows = [
        ("Score-versus-return rank correlation", num(original_summary["spearman_score_vs_usd_return"]), num(enhanced_summary["spearman_score_vs_usd_return"])),
        ("Top-minus-bottom USD return spread", pct(original_summary["top_minus_bottom_spread"]), pct(enhanced_summary["top_minus_bottom_spread"])),
        ("Directional hit rate", pct(original_summary["directional_hit_rate"]), pct(enhanced_summary["directional_hit_rate"])),
        ("Buy / Watch / Reject", f"{original_summary['buy_candidates']} / {original_summary['watch_companies']} / {original_summary['reject_companies']}", f"{enhanced_summary['buy_candidates']} / {enhanced_summary['watch_companies']} / {enhanced_summary['reject_companies']}"),
    ]
    add_table(doc, ["Evaluation measure", "Original", "Cycle-aware"], comparison_rows, widths=[3.35, 1.55, 1.55], font_size=9.2)
    doc.add_paragraph(
        "The cycle-aware version fixed the most important decision error: it no longer rejected any of the five memory producers, and its memory-only directional hit rate was 100%. It also changed the six-company top-minus-bottom spread from –29.5 to +35.4 percentage points. It did not establish reliable fine ranking: overall Spearman correlation remained –0.20, and correlation inside the five-company memory group was only +0.10."
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(semiconductor_chart_path), width=Inches(6.75))
    caption = doc.add_paragraph("Figure 3  Original long-term semiconductor score versus the cycle-aware score. TSMC remains on the long-term scorecard.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].font.italic = True
    caption.runs[0].font.size = Pt(9)

    doc.add_page_break()
    doc.add_heading("Company-level semiconductor results", level=1)
    targeted_rows = []
    original_by_ticker = {row["ticker"]: row for row in initial_semiconductor["records"]}
    for record in semiconductor["records"]:
        targeted_rows.append([
            record["ticker"],
            record["company"],
            f"{original_by_ticker[record['ticker']]['score']:.1f}",
            f"{record['score']:.1f}",
            record["verdict"].replace(" Candidate", ""),
            pct(record["usd_return"]),
            pct(record["usd_excess_return"]),
        ])
    add_table(
        doc,
        ["Ticker", "Company", "Original", "Enhanced", "Verdict", "USD return", "USD excess"],
        targeted_rows,
        widths=[0.72, 1.75, 0.70, 0.70, 0.82, 0.82, 0.82],
        font_size=7.7,
    )
    doc.add_heading("Two scorecards for two economic structures", level=1)
    scorecard_rows = [
        ("Long-term semiconductor", "Foundries, fabless designers, equipment and diversified chip firms", "Quality 25%; growth/cycle 25%; valuation 15%; balance sheet 15%; innovation/reinvestment 10%; risk/momentum 10%"),
        ("Tactical memory semiconductor", "DRAM, HBM and memory manufacturers with pronounced pricing cycles", "Quality 15%; cycle/earnings inflection 25%; memory market cycle 20%; valuation 10%; balance sheet 10%; innovation 5%; risk/momentum 15%"),
    ]
    add_table(doc, ["Scorecard", "Use", "Pillar weights"], scorecard_rows, widths=[1.55, 2.35, 2.75], font_size=8.6)
    doc.add_paragraph(
        "The memory-semiconductor model is fully automated. It measures reported revenue acceleration, gross-margin change, inventory-days normalization, cash generation, valuation, R&D intensity, balance-sheet resilience, market risk, and relative momentum without an analyst-supplied cycle input."
    )
    doc.add_heading("Interpretation guardrail", level=2)
    doc.add_paragraph(
        "This redesign was informed by the same six-month outcome used to evaluate it. The improvement is therefore in-sample and must not be presented as independent validation. It is a better-specified hypothesis for the next walk-forward test. The weights and anchors should now be frozen before testing earlier and later memory cycles."
    )

    doc.add_heading("Limitations", level=1)
    limitations = [
        "The sample contains 18 companies and one six-month market regime. Statistical conclusions are fragile.",
        "Only one company qualified as a Buy Candidate, so Buy performance cannot be generalized.",
        "Yahoo Finance statement periods were filtered to dates on or before the scoring date, but historical filing publication timestamps and restatement vintages were unavailable. The test therefore cannot guarantee a fully point-in-time financial dataset.",
        "The universe was selected to cover markets and team mandates rather than through a randomized or exhaustive sampling procedure.",
        "The test evaluates rank and verdict outcomes, not portfolio position sizing, turnover, transaction costs, taxes, liquidity, or the defensive sleeve.",
        "A six-month realization is shorter than the model's stated three-to-five-year holding horizon.",
        "The cycle-aware memory scorecard was designed after observing this cohort's outcome, so its improved classifications are in-sample and subject to overfitting.",
    ]
    for item in limitations:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Recommended next steps", level=1)
    recommendations = [
        ("Audit automated data gaps", "Review confidence and missing-field warnings before comparing banks, insurers, and operating companies."),
        ("Build a true point-in-time feed", "Use SEC and local-regulator filing timestamps so each historical run includes only information publicly available on that date."),
        ("Expand the time series", "Repeat the test at monthly or quarterly starting dates over at least five years and evaluate rolling one-, three-, six-, and twelve-month outcomes."),
        ("Use sector-aware comparisons", "Compare ranks within the same scorecard and sector before interpreting absolute score differences."),
        ("Pre-register evaluation rules", "Choose the universe, horizons, benchmark, and success metrics before running the next test to reduce selection bias."),
        ("Validate the memory scorecard", "Freeze the new specification and test it across at least two prior DRAM upturns, one downturn, and future unseen periods."),
    ]
    add_table(doc, ["Priority", "Action"], recommendations, widths=[1.75, 5.05], font_size=9.4)

    doc.add_heading("Replication", level=1)
    doc.add_paragraph(
        "Install the repository using install_windows.bat, then run the command below from the repository folder. The included universe file contains all 18 tickers. Frozen input snapshots and the resulting CSV and JSON files accompany this report for auditability."
    )
    command = doc.add_paragraph()
    command.paragraph_format.left_indent = Inches(0.25)
    run = command.add_run("wharton backtest --start 2026-03-20 --end 2026-09-18")
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    doc.add_paragraph(
        "To reproduce the score from an existing snapshot without downloading new scoring inputs, place the frozen snapshots in the selected output folder and add --reuse-snapshots. Market-data providers can revise historical datasets, so the frozen results should remain the audit record."
    )
    enhanced_command = doc.add_paragraph()
    enhanced_command.paragraph_format.left_indent = Inches(0.25)
    run = enhanced_command.add_run("wharton backtest --universe config/enhanced_backtest_universe.csv --start 2026-03-20 --end 2026-09-18")
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    memory_command = doc.add_paragraph()
    memory_command.paragraph_format.left_indent = Inches(0.25)
    run = memory_command.add_run("wharton backtest --universe config/semiconductor_backtest_universe.csv --start 2026-03-20 --end 2026-09-18")
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "The broad test supports keeping the model as a disciplined research framework, and the semiconductor diagnostic shows why cyclical industries need structurally appropriate factors. Neither result proves reliable return forecasting. The team's best next move is to freeze this specification, add true point-in-time filing data and verified specialist metrics, and repeat the tests across many unseen dates before using scores as a central portfolio decision rule."
    )

    doc.add_heading("Data sources", level=1)
    doc.add_paragraph("Yahoo Finance historical prices, adjusted prices, corporate actions, benchmark series, FX series, and available financial statements: https://finance.yahoo.com")
    doc.add_paragraph("Yahoo Finance company profiles, statements, prices, benchmarks, exchange rates, and timestamped news were used for automated scoring inputs.")
    doc.add_paragraph("Micron Technology investor relations, 18 March 2026, fiscal Q2 2026 prepared remarks used as qualitative cycle-context validation: https://investors.micron.com/static-files/e089f8c0-065d-47b8-9d02-bfa863cdb357")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Growth Sleeve Scoring Model Backtest and Semiconductor Cycle Analysis"
    doc.core_properties.subject = "Retrospective model evaluation and replication guide"
    doc.core_properties.author = "Wharton Global High School Investment Competition Team"
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 7:
        raise SystemExit("Usage: build_backtest_report.py BASELINE_JSON ENHANCED_JSON SEMICONDUCTOR_JSON INITIAL_SEMICONDUCTOR_JSON OUTPUT_DOCX CHART_DIR")
    build_report(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]), Path(sys.argv[4]), Path(sys.argv[5]), Path(sys.argv[6]))
