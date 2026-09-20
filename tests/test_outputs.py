import hashlib
import sqlite3

from openpyxl import load_workbook

from growth_scorer.engine import score_snapshot
from growth_scorer.reporting import write_workbook
from growth_scorer.storage import record_and_rank


def test_workbook_has_required_sheets_and_is_reproducible(tmp_path, snapshot_factory):
    snapshot = snapshot_factory("general", 2)
    result = score_snapshot(snapshot)
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    write_workbook(first, result, snapshot)
    write_workbook(second, result, snapshot)
    assert load_workbook(first, read_only=True).sheetnames == [
        "Summary",
        "Factor Breakdown",
        "Raw Data",
        "Source Log",
        "Warnings",
    ]
    assert hashlib.sha256(first.read_bytes()).digest() == hashlib.sha256(second.read_bytes()).digest()


def test_rank_requires_two_comparable_companies(tmp_path, snapshot_factory):
    database = tmp_path / "history.sqlite3"
    first = score_snapshot(snapshot_factory("general", 2, ticker="AAA"))
    rank, count = record_and_rank(database, first)
    assert rank is None and count == 1
    second = score_snapshot(snapshot_factory("general", 1, ticker="BBB"))
    rank, count = record_and_rank(database, second)
    assert rank == 2 and count == 2
    with sqlite3.connect(database) as connection:
        assert connection.execute("SELECT COUNT(*) FROM score_history").fetchone()[0] == 2
