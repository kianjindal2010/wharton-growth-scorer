from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from growth_scorer.batch import BatchItem, read_batch_csv, run_batch
from growth_scorer.cli import build_parser


def test_batch_parser_accepts_csv_mode(tmp_path):
    batch_file = tmp_path / "stocks.csv"
    args = build_parser().parse_args([
        "batch", "--input", str(batch_file), "--as-of", "2026-09-20",
    ])
    assert args.command == "batch"
    assert args.input == batch_file
    assert args.as_of == date(2026, 9, 20)


def test_read_batch_csv_supports_optional_scorecard(tmp_path):
    path = tmp_path / "stocks.csv"
    path.write_text(
        "ticker,country,scorecard\n"
        "msft,us,auto\n"
        "2330.tw,tw,semiconductor\n",
        encoding="utf-8",
    )
    items = read_batch_csv(path)
    assert items[0] == BatchItem("MSFT", "US", "auto")
    assert items[1].ticker == "2330.TW"


def test_batch_creates_summary_and_individual_reports(tmp_path, snapshot_factory, monkeypatch):
    def fake_build(ticker, country, as_of):
        anchor = 2 if ticker == "AAA" else 1
        return snapshot_factory("general", anchor, ticker=ticker, country=country, as_of=as_of)

    monkeypatch.setattr("growth_scorer.batch.build_snapshot", fake_build)
    workbook_path, json_path, payload = run_batch(
        [BatchItem("AAA", "US"), BatchItem("BBB", "US")],
        date(2026, 9, 20),
        tmp_path / "output",
        tmp_path / "snapshots",
        tmp_path / "history.sqlite3",
    )

    assert workbook_path.exists()
    assert json_path.exists()
    assert payload["successful"] == 2
    assert payload["failed"] == 0
    assert [record["rank"] for record in payload["results"]] == [1, 2]
    assert all(Path(record["workbook_path"]).exists() for record in payload["results"])
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == ["Consolidated", "Batch Summary", "Results", "Failures", "Instructions"]
    consolidated = workbook["Consolidated"]
    assert consolidated.freeze_panes == "A2"
    assert consolidated.auto_filter.ref == "A1:F3"
    assert [cell.value for cell in consolidated[1]] == [
        "Stock", "ticker", "Industry", "Country Code", "Model Score", "Model Verdict",
    ]
    rows = list(consolidated.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert rows[0][0] == "Synthetic Company"
    assert rows[0][1] == "AAA"
    assert rows[0][2] == "Conglomerates"
    assert rows[0][3] == "US"
    assert isinstance(rows[0][4], (int, float))
    assert rows[0][5] in {"Buy Candidate", "Watch", "Reject", "Insufficient Data"}


def test_batch_failure_does_not_stop_valid_companies(tmp_path, snapshot_factory, monkeypatch):
    monkeypatch.setattr(
        "growth_scorer.batch.build_snapshot",
        lambda ticker, country, as_of: snapshot_factory("general", 2, ticker=ticker, country=country, as_of=as_of),
    )
    workbook_path, _, payload = run_batch(
        [BatchItem("AAA", "US"), BatchItem("2330.TW", "US")],
        date(2026, 9, 20),
        tmp_path / "output",
        tmp_path / "snapshots",
        tmp_path / "history.sqlite3",
    )
    assert payload["successful"] == 1
    assert payload["failed"] == 1
    assert payload["failures"][0]["ticker"] == "2330.TW"
    workbook = load_workbook(workbook_path, read_only=True)
    rows = list(workbook["Consolidated"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    failed_row = next(row for row in rows if row[1] == "2330.TW")
    assert failed_row[4] is None
    assert failed_row[5] == "Failed"
