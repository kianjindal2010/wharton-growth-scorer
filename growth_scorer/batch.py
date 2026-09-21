from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import build_snapshot
from .engine import score_snapshot
from .overrides import apply_overrides, read_overrides
from .reporting import write_json, write_workbook
from .storage import record_and_rank
from .validation import SCORECARDS, safe_name, validate_ticker_country


NAVY = "17365D"
BLUE = "2079B0"
LIGHT_BLUE = "D9EAF7"
RED = "BE1E2D"
WHITE = "FFFFFF"


@dataclass(frozen=True)
class BatchItem:
    ticker: str
    country: str
    scorecard: str = "auto"
    overrides: Path | None = None

    def normalized(self) -> "BatchItem":
        return BatchItem(
            ticker=self.ticker.strip().upper(),
            country=self.country.strip().upper(),
            scorecard=(self.scorecard or "auto").strip().lower(),
            overrides=self.overrides,
        )


def read_batch_csv(path: Path) -> list[BatchItem]:
    if not path.exists():
        raise FileNotFoundError(f"Batch input file not found: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = {str(name).strip().lower() for name in (reader.fieldnames or [])}
        if not {"ticker", "country"}.issubset(headers):
            raise ValueError("Batch CSV must contain ticker and country columns")
        items: list[BatchItem] = []
        for row_number, raw in enumerate(reader, start=2):
            row = {str(key).strip().lower(): (value or "").strip() for key, value in raw.items()}
            if not any(row.values()):
                continue
            if not row.get("ticker") or not row.get("country"):
                raise ValueError(f"Batch CSV row {row_number} must contain ticker and country")
            override_value = row.get("overrides", "").strip('"')
            override_path = Path(override_value) if override_value else None
            if override_path and not override_path.is_absolute():
                override_path = path.parent / override_path
            items.append(BatchItem(
                ticker=row["ticker"],
                country=row["country"],
                scorecard=row.get("scorecard") or "auto",
                overrides=override_path,
            ).normalized())
    if not items:
        raise ValueError("Batch CSV contains no companies")
    return items


def _validate_items(items: list[BatchItem]) -> list[BatchItem]:
    if not items:
        raise ValueError("Add at least one company to the batch")
    normalized: list[BatchItem] = []
    seen: set[str] = set()
    for item in items:
        item = item.normalized()
        if item.ticker in seen:
            raise ValueError(f"Duplicate batch ticker: {item.ticker}")
        seen.add(item.ticker)
        normalized.append(item)
    return normalized


def _fit(ws, maximum: int = 70) -> None:
    for column_cells in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, maximum)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 11)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _write_batch_workbook(path: Path, payload: dict[str, Any]) -> None:
    workbook = Workbook()
    workbook.properties.creator = "Wharton Growth Scorer"
    ws = workbook.active
    ws.title = "Batch Summary"
    ws.append(["Wharton Growth Scorer Batch", "Result"])
    summary_rows = [
        ("As-of date", payload["as_of"]),
        ("Started", payload["started_at"]),
        ("Completed", payload["completed_at"]),
        ("Requested companies", payload["requested"]),
        ("Successfully scored", payload["successful"]),
        ("Failed", payload["failed"]),
        ("Buy Candidates", payload["verdict_counts"].get("Buy Candidate", 0)),
        ("Watch", payload["verdict_counts"].get("Watch", 0)),
        ("Reject", payload["verdict_counts"].get("Reject", 0)),
        ("Insufficient Data", payload["verdict_counts"].get("Insufficient Data", 0)),
        ("Batch directory", payload["batch_directory"]),
    ]
    for row in summary_rows:
        ws.append(row)
    _header(ws)
    _fit(ws, 100)
    ws.auto_filter.ref = "A1:B1"

    results = workbook.create_sheet("Results")
    results.append([
        "Ticker", "Company", "Country", "Scorecard", "Score", "Verdict", "Confidence",
        "Batch rank", "Comparable count", "Risk gates", "Warnings", "Excel report", "JSON result",
    ])
    for record in payload["results"]:
        results.append([
            record["ticker"], record["company_name"], record["country"], record["scorecard"],
            record["score"], record["verdict"], record["confidence"] / 100,
            record["rank"] if record["rank"] is not None else "Not meaningful",
            record["comparable_count"], "; ".join(record["risk_gates"]),
            "; ".join(record["warnings"]), "Open company workbook", "Open JSON result",
        ])
        row = results.max_row
        results.cell(row, 7).number_format = "0.0%"
        for column, key in ((12, "workbook_path"), (13, "json_path")):
            linked_path = Path(record[key])
            results.cell(row, column).hyperlink = linked_path.resolve().as_uri()
            results.cell(row, column).style = "Hyperlink"
        verdict_cell = results.cell(row, 6)
        if record["verdict"] in {"Reject", "Insufficient Data"}:
            verdict_cell.font = Font(color=RED, bold=True)
    _header(results)
    _fit(results, 90)

    failures = workbook.create_sheet("Failures")
    failures.append(["Ticker", "Country", "Scorecard", "Error"])
    for failure in payload["failures"]:
        failures.append([failure["ticker"], failure["country"], failure["scorecard"], failure["error"]])
    if failures.max_row == 1:
        failures.append(["", "", "", "No failures"])
    _header(failures)
    _fit(failures, 100)

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["How to use this workbook", "Explanation"])
    guidance = [
        ("Comparison", "Rank companies only inside the same scorecard; different scorecards are not directly comparable."),
        ("Rank", "A rank is meaningful only when at least two successful companies use the same scorecard."),
        ("Detailed evidence", "Open each company workbook for factor scores, raw data, sources, warnings, and risk gates."),
        ("Failures", "A failed company does not stop the rest of the batch. Correct its ticker, country, data, or override and rerun it."),
        ("Decision rule", "This is a research consistency tool, not a return forecast or instruction to trade."),
    ]
    for row in guidance:
        instructions.append(row)
    _header(instructions)
    _fit(instructions, 100)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def run_batch(
    items: list[BatchItem],
    as_of: date,
    output_dir: Path,
    data_dir: Path,
    history_db: Path,
) -> tuple[Path, Path, dict[str, Any]]:
    items = _validate_items(items)
    started = datetime.now().astimezone()
    stamp = started.strftime("%Y%m%d_%H%M%S")
    batch_dir = output_dir / as_of.isoformat() / f"batch_{stamp}"
    company_dir = batch_dir / "companies"
    successes: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []

    for position, item in enumerate(items, start=1):
        print(f"[{position}/{len(items)}] Scoring {item.ticker} ({item.country})...")
        try:
            validate_ticker_country(item.ticker, item.country)
            if item.scorecard not in SCORECARDS:
                raise ValueError(f"Unknown scorecard: {item.scorecard}")
            snapshot = build_snapshot(item.ticker, item.country, as_of)
            if item.overrides:
                records, warnings = read_overrides(item.overrides, as_of)
                apply_overrides(snapshot, records, warnings)
            ticker_name = safe_name(snapshot.ticker)
            frozen_path = data_dir / ticker_name / f"{as_of.isoformat()}.json"
            snapshot.save(frozen_path)
            result = score_snapshot(snapshot, item.scorecard)
            result.snapshot_path = str(frozen_path.resolve())
            successes.append({"item": item, "snapshot": snapshot, "result": result})
        except Exception as exc:
            failures.append({
                "ticker": item.ticker, "country": item.country,
                "scorecard": item.scorecard, "error": str(exc),
            })
            print(f"  Could not score {item.ticker}: {exc}")

    by_scorecard: dict[str, list[dict[str, Any]]] = {}
    for success in successes:
        by_scorecard.setdefault(success["result"].scorecard, []).append(success)
    for group in by_scorecard.values():
        group.sort(key=lambda entry: (-entry["result"].score, entry["result"].ticker))
        count = len(group)
        for rank, success in enumerate(group, start=1):
            result = success["result"]
            result.comparable_count = count
            result.rank = rank if count >= 2 else None
            if count < 2:
                result.warnings.append(
                    "Batch rank is not meaningful until at least two companies use the same scorecard"
                )

    result_records: list[dict[str, Any]] = []
    for success in successes:
        snapshot = success["snapshot"]
        result = success["result"]
        base = f"{safe_name(result.ticker)}_{as_of.isoformat()}_{result.scorecard}"
        workbook_path = company_dir / f"{base}.xlsx"
        json_path = company_dir / f"{base}.json"
        result.workbook_path = str(workbook_path.resolve())
        result.json_path = str(json_path.resolve())
        write_workbook(workbook_path, result, snapshot)
        write_json(json_path, result)
        record_and_rank(history_db, result)
        result_records.append({
            "ticker": result.ticker,
            "company_name": result.company_name or "",
            "country": result.country,
            "scorecard": result.scorecard,
            "score": result.score,
            "verdict": result.verdict,
            "confidence": result.confidence,
            "rank": result.rank,
            "comparable_count": result.comparable_count,
            "risk_gates": result.risk_gates,
            "warnings": result.warnings,
            "snapshot_path": result.snapshot_path,
            "workbook_path": result.workbook_path,
            "json_path": result.json_path,
        })

    result_records.sort(key=lambda record: (record["scorecard"], -record["score"], record["ticker"]))
    verdict_counts: dict[str, int] = {}
    for record in result_records:
        verdict_counts[record["verdict"]] = verdict_counts.get(record["verdict"], 0) + 1
    completed = datetime.now().astimezone()
    payload: dict[str, Any] = {
        "as_of": as_of.isoformat(),
        "started_at": started.isoformat(),
        "completed_at": completed.isoformat(),
        "requested": len(items),
        "successful": len(result_records),
        "failed": len(failures),
        "verdict_counts": verdict_counts,
        "batch_directory": str(batch_dir.resolve()),
        "results": result_records,
        "failures": failures,
    }
    summary_workbook = batch_dir / "Batch_Summary.xlsx"
    summary_json = batch_dir / "Batch_Summary.json"
    _write_batch_workbook(summary_workbook, payload)
    summary_json.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return summary_workbook, summary_json, payload
