from __future__ import annotations

import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import InputSnapshot, ScoreResult


NAVY = "17365D"
BLUE = "2079B0"
LIGHT_BLUE = "D9EAF7"
RED = "BE1E2D"
WHITE = "FFFFFF"
GREEN = "DDEBF7"


def _header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _fit(ws, maximum: int = 60) -> None:
    for column_cells in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, maximum)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 11)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_summary(workbook: Workbook, result: ScoreResult, snapshot: InputSnapshot) -> None:
    ws = workbook.active
    ws.title = "Summary"
    ws.append(["Growth Sleeve Scoring Model", "Result"])
    fields = [
        ("Ticker", result.ticker),
        ("Company", result.company_name or ""),
        ("Country", result.country),
        ("Yahoo sector", snapshot.sector or ""),
        ("Yahoo industry", snapshot.industry or ""),
        ("Yahoo industry key", snapshot.industry_key or ""),
        ("As-of date", result.as_of.isoformat()),
        ("Scorecard", result.scorecard.title()),
        ("Classification confidence", result.classification_confidence),
        ("Classification reason", result.classification_reason),
        ("Overall score", result.score),
        ("Verdict", result.verdict),
        ("Data confidence", result.confidence / 100),
        ("Comparable rank", "Not meaningful" if result.rank is None else f"{result.rank} of {result.comparable_count}"),
        ("Growth sleeve policy", "50% of total portfolio; initial reference amount $150,000"),
        ("Last price date", snapshot.last_price_date.isoformat() if snapshot.last_price_date else ""),
        ("Latest financial period", snapshot.latest_financial_date.isoformat() if snapshot.latest_financial_date else ""),
        ("Risk gates", "; ".join(result.risk_gates) if result.risk_gates else "None"),
    ]
    field_rows: dict[str, int] = {}
    for field, value in fields:
        ws.append([field, value])
        field_rows[field] = ws.max_row
    ws.append([])
    ws.append(["Pillar", "Score (0-100)"])
    for pillar, score in result.pillar_scores.items():
        ws.append([pillar, score])
    _header(ws, 1)
    _header(ws, len(fields) + 3)
    ws.cell(field_rows["Verdict"], 2).font = Font(
        bold=True, color=RED if result.verdict in {"Reject", "Insufficient Data"} else NAVY
    )
    ws.cell(field_rows["Data confidence"], 2).number_format = "0.0%"
    _fit(ws)
    ws.auto_filter.ref = "A1:B1"


def _write_factors(workbook: Workbook, result: ScoreResult) -> None:
    ws = workbook.create_sheet("Factor Breakdown")
    ws.append([
        "Pillar", "Metric", "Metric key", "Weight", "Raw value", "Unit", "Score",
        "Contribution", "Observed", "Source", "Bad anchor", "Neutral anchor", "Excellent anchor",
    ])
    for factor in result.factors:
        ws.append([
            factor.pillar,
            factor.label,
            factor.metric,
            factor.weight / 100,
            factor.raw_value,
            factor.unit,
            factor.score,
            factor.contribution,
            "Yes" if factor.observed else "No",
            factor.source,
            factor.bad_anchor,
            factor.neutral_anchor,
            factor.excellent_anchor,
        ])
    _header(ws)
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 4).number_format = "0.0%"
        if ws.cell(row, 9).value == "No":
            for column in range(1, ws.max_column + 1):
                ws.cell(row, column).fill = PatternFill("solid", fgColor="FFF2CC")
    _fit(ws)


def _flatten_raw(snapshot: InputSnapshot) -> list[list[Any]]:
    rows: list[list[Any]] = [
        ["company_profile", snapshot.as_of.isoformat(), "sector", snapshot.sector],
        ["company_profile", snapshot.as_of.isoformat(), "sector_key", snapshot.sector_key],
        ["company_profile", snapshot.as_of.isoformat(), "industry", snapshot.industry],
        ["company_profile", snapshot.as_of.isoformat(), "industry_key", snapshot.industry_key],
        ["company_profile", snapshot.as_of.isoformat(), "quote_type", snapshot.quote_type],
        ["company_profile", snapshot.as_of.isoformat(), "exchange", snapshot.exchange],
        ["company_profile", snapshot.as_of.isoformat(), "business_summary", snapshot.business_summary],
    ]
    for metric, value in sorted(snapshot.metrics.items()):
        rows.append(["calculated_metrics", snapshot.as_of.isoformat(), metric, value])
    datasets = {
        "prices": snapshot.prices,
        "benchmark_prices": snapshot.benchmark_prices,
        "fx_rates": snapshot.fx_rates,
        "news": snapshot.news,
    }
    datasets.update(snapshot.statements)
    for dataset, records in datasets.items():
        for record in records:
            period = record.get("date") or record.get("period") or record.get("published_at")
            for field, value in record.items():
                if field not in {"date", "period", "published_at"}:
                    rows.append([dataset, period, field, value])
    return rows


def _write_raw(workbook: Workbook, snapshot: InputSnapshot) -> None:
    ws = workbook.create_sheet("Raw Data")
    ws.append(["Dataset", "Date / period", "Field", "Value"])
    for row in _flatten_raw(snapshot):
        ws.append(row)
    _header(ws)
    _fit(ws)


def _write_sources(workbook: Workbook, snapshot: InputSnapshot) -> None:
    ws = workbook.create_sheet("Source Log")
    ws.append(["Source", "Retrieval time", "Reporting period", "URL", "Override provenance"])
    for source in snapshot.source_log:
        ws.append([
            source.source,
            source.retrieval_time.isoformat(),
            source.reporting_period,
            source.url,
            source.provenance,
        ])
    _header(ws)
    _fit(ws)


def _write_warnings(workbook: Workbook, result: ScoreResult) -> None:
    ws = workbook.create_sheet("Warnings")
    ws.append(["Type", "Message"])
    for gate in result.risk_gates:
        ws.append(["Risk gate", gate])
    for warning in result.warnings:
        ws.append(["Warning", warning])
    if ws.max_row == 1:
        ws.append(["Information", "No warnings or risk gates"])
    _header(ws)
    _fit(ws, maximum=100)


def _normalize_xlsx(path: Path) -> None:
    """Normalize ZIP member timestamps so frozen inputs produce stable workbooks."""
    with NamedTemporaryFile(suffix=".xlsx", delete=False, dir=path.parent) as temp:
        temporary = Path(temp.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = source.getinfo(name).external_attr
                payload = source.read(name)
                if name == "docProps/core.xml":
                    payload = re.sub(
                        rb"(<dcterms:modified[^>]*>)[^<]+(</dcterms:modified>)",
                        rb"\g<1>2020-01-01T00:00:00Z\g<2>",
                        payload,
                    )
                target.writestr(info, payload)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def write_workbook(path: Path, result: ScoreResult, snapshot: InputSnapshot) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    fixed_time = datetime.combine(snapshot.as_of, datetime.min.time())
    workbook.properties.creator = "Growth Sleeve Scoring Model"
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    _write_summary(workbook, result, snapshot)
    _write_factors(workbook, result)
    _write_raw(workbook, snapshot)
    _write_sources(workbook, snapshot)
    _write_warnings(workbook, result)
    workbook.save(path)
    _normalize_xlsx(path)


def write_json(path: Path, result: ScoreResult) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(result.model_dump_json(indent=2), encoding="utf-8")
